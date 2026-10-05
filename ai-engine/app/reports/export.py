"""Export a full report to JSON / CSV / XLSX / PDF (all local, no external service)."""
from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from typing import Any


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:g}"
    if isinstance(v, (list, tuple)):
        return ", ".join(_fmt(x) for x in v)
    return str(v)


def tables(r: dict) -> dict[str, tuple[list[str], list[list]]]:
    """Tabular sections: name -> (headers, rows). Includes evidence message ids for traceability."""
    t = lambda ts: [[x["title"], x["assigned_to"], x["status"], x["priority"], x["weight"], x["deadline"], x["completed_at"],  # noqa: E731
                     x["confidence"], x["needs_review"], x["evidence_message_ids"]] for x in ts]
    th = ["Task", "Assigned to", "Status", "Priority", "Weight", "Deadline", "Completed at", "Confidence", "Needs review", "Source messages"]
    all_tasks = r["9_completed_work"] + r["10_in_progress_work"] + r["11_pending_work"] + r["12_blocked_work"]
    return {
        "Tasks": (th, t(all_tasks)),
        "Client pending": (["Item", "Type", "Requested by", "Waiting since", "Age (days)", "Priority", "Impact", "Status", "Confidence", "Source messages"],
                           [[c["title"], c["type"], c["requested_by"], c["waiting_since"], c["age_days"], c["priority"], c["impact"],
                             c["status"], c["confidence"], c["source_message_ids"]] for c in r["13_client_pending"]]),
        "Team performance": (["Person", "Role", "Assigned", "Completed", "Overdue", "Weighted completion %", "On-time %", "Score", "Low sample"],
                             [[c["person"], c["role"] or "unknown", c["assigned"], c["completed"], c["overdue"], c["weighted_completion_pct"],
                               c["on_time_pct"], c["score"], c["low_sample"]] for c in r["25_team_performance"]]),
        "Risks": (["Risk", "Severity", "Description", "Recommended action", "Source messages"],
                  [[k["title"], k["severity"], k["description"], k["recommended_action"], k["source_message_ids"]] for k in r["22_risks"]["risks"]]),
        "Bugs": (["Bug", "Status", "Reported by", "Reported at", "Source messages"],
                 [[b["title"], b["status"], b["reported_by"], b["reported_at"], b["source_message_ids"]] for b in r["15_bugs"]]),
        "Requirement changes": (["Requirement", "Old", "New", "Changed by", "At", "Confidence", "Needs review", "Source messages"],
                                [[c["key"], c["old_value"], c["new_value"], c["changed_by"], c["changed_at"], c["confidence"],
                                  c["needs_review"], c["source_message_ids"]] for c in r["17_requirement_changes"]]),
        "Conflicts": (["Topic", "Statement A", "Author A", "Statement B", "Author B", "Severity", "Status"],
                      [[c["topic"], c["statement_a"]["text"], c["statement_a"]["author"], c["statement_b"]["text"],
                        c["statement_b"]["author"], c["severity"], c["status"]] for c in r["23b_conflicts"]]),
        "Timeline": (["When", "Type", "Event", "Source messages"],
                     [[e["at"], e["type"], e["title"], e["source_message_ids"]] for e in r["24_timeline"]]),
        "Next actions": (["Action", "Owner", "Priority", "Deadline", "Reason", "Source messages"],
                         [[a["action"], a["owner"], a["priority"], a["deadline"], a["reason"], a["source_message_ids"]] for a in r["28_next_actions"]]),
    }


def to_json(r: dict) -> str:
    return json.dumps(r, default=_fmt, ensure_ascii=False, indent=2)


def to_csv(r: dict, section: str = "Tasks") -> str:
    headers, rows = tables(r)[section]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for row in rows:
        w.writerow([_fmt(c) for c in row])
    return buf.getvalue()


def to_xlsx(r: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    meta, P, H = r["meta"], r["3_progress"], r["4_health"]
    rows = [("Project", meta["project"]), ("Client", meta["client"]), ("Company", meta["company"]),
            ("Period", f"{_fmt(meta['reporting_period'][0])} to {_fmt(meta['reporting_period'][1])}"),
            ("Overall progress %", P["overall"]), ("Health", f"{H['score']} ({H['status']})"), ("", ""),
            ("Executive summary", "")] + [("", c["text"]) for c in r["1_executive_summary"]["claims"]]
    for row in rows:
        ws.append([_fmt(c) for c in row])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 110
    for c in ws["A"]:
        c.font = Font(bold=True)
    for name, (headers, data) in tables(r).items():
        s = wb.create_sheet(name[:31])
        s.append(headers)
        for cell in s[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="3B5BDB")
        for row in data:
            s.append([_fmt(c) for c in row])
        for i, h in enumerate(headers, 1):
            width = max([len(h)] + [len(_fmt(row[i - 1])) for row in data[:200]])
            s.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
            for c in s[get_column_letter(i)][1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(r: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    ss = getSampleStyleSheet()
    cell = ss["BodyText"].clone("cell", fontSize=7.5, leading=9)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"Project report - {r['meta']['project']}")
    meta, P, H = r["meta"], r["3_progress"], r["4_health"]
    esc = lambda s: _fmt(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    el = [Paragraph(esc(meta["company"] or "Project report"), ss["Title"]),
          Paragraph(f"<b>{esc(meta['project'])}</b> &mdash; client: {esc(meta['client'])}", ss["Heading3"]),
          Paragraph(f"Period: {esc(meta['reporting_period'][0])} to {esc(meta['reporting_period'][1])} &nbsp; Generated: {esc(meta['generated_at'])}", ss["Normal"]),
          Spacer(1, 6)]

    def kpis(items):
        t = Table([[Paragraph(f"<font size=15><b>{esc(v)}</b></font><br/><font size=7>{esc(k)}</font>", ss["Normal"]) for k, v in items]],
                  colWidths=[(180 / len(items)) * mm] * len(items))
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey), ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                               ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F6FB")), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
        return t

    el.append(kpis([("Progress", f"{P['overall']}%" if P["overall"] is not None else "n/a"), (f"Health ({H['status'].replace('_', ' ')})", f"{H['score']:g}"),
                    ("Completed", len(r["9_completed_work"])), ("Pending", len(r["11_pending_work"]) + len(r["10_in_progress_work"])),
                    ("Blocked", len(r["12_blocked_work"])), ("Client pending", r["8_client_activity"]["open_items"]),
                    ("Overdue", len(r["21_overdue_tasks"]))]))
    el += [Spacer(1, 8), Paragraph("Executive summary", ss["Heading2"])]
    el += [Paragraph(esc(c["text"]), ss["BodyText"]) for c in r["1_executive_summary"]["claims"]]
    sub = [(k.replace("_", " ").title(), f"{P[k]:g}%") for k in ("development", "design", "testing", "deployment", "client_approval") if P.get(k) is not None]
    if sub:
        el += [Spacer(1, 6), kpis(sub)]
    for name, (headers, rows) in tables(r).items():
        if not rows:
            continue
        keep = [i for i, h in enumerate(headers) if h != "Source messages"][:7]
        data = [[Paragraph(f"<b>{esc(headers[i])}</b>", cell) for i in keep]] + \
               [[Paragraph(esc(row[i])[:140], cell) for i in keep] for row in rows[:60]]
        t = Table(data, repeatRows=1, colWidths=[(182 / len(keep)) * mm] * len(keep))
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7ECFB")), ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        el += [Spacer(1, 10), Paragraph(esc(name), ss["Heading2"]), t]
        if len(rows) > 60:
            el.append(Paragraph(f"... {len(rows) - 60} more rows in the Excel / JSON export.", cell))
    if r["27_recommendations"]:
        el += [Spacer(1, 10), Paragraph("Recommendations", ss["Heading2"])]
        el += [Paragraph("&bull; " + esc(x["recommendation"]) + f" <i>({esc(x['because'])})</i>", ss["BodyText"]) for x in r["27_recommendations"]]
    el += [Spacer(1, 10), Paragraph("Every item is traceable to the original WhatsApp messages (message ids in the Excel/JSON export). "
                                    "Figures are AI-extracted and may need review.", cell)]
    doc.build(el)
    return buf.getvalue()
