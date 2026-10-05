import io
import json
from datetime import date

import pytest

from app.evaluation.rouge import rouge
from app.intelligence.pipeline import analyze
from app.intelligence.types import Msg
from app.parsers.whatsapp import WhatsAppParser
from app.reports import export
from app.reports.builder import build_timeline, daily_report, full_report, week_bounds, weekly_report
from app.summarization.executive import build_summary, verify_summary
from app.summarization.extractive import informative, summarize_messages
from tests.scenario_chat import CHAT

META = {"company": "Acme Tech", "name": "E-Commerce Website", "client": "Sharma Retail", "manager": "Rahul"}


def parse(chat):
    return [Msg.from_parsed(m) for m in WhatsAppParser().parse_text(chat)]


@pytest.fixture(scope="module")
def msgs():
    return parse(CHAT)


@pytest.fixture(scope="module")
def result(msgs):
    return analyze(msgs)


@pytest.fixture(scope="module")
def report(result, msgs):
    return full_report(result, msgs, META)


# --------------------------------------------------------------- extractive
def test_informative_prior_filters_chatter():
    m = lambda t: type("M", (), {"sender": "A", "text": t})()  # noqa: E731
    assert informative(m("ok")) == 0 and informative(m("thanks!")) == 0 and informative(m("👍")) == 0
    assert informative(m("Razorpay API integrated and tested successfully")) > 0.4


def test_extractive_returns_original_messages_only(msgs):
    out = summarize_messages(msgs, k=5)
    by_id = {m.id: m for m in msgs}
    assert 0 < len(out) <= 5
    assert all(by_id[o["message_id"]].text == o["text"] for o in out)  # verbatim quotes, never rewritten
    assert [o["ts"] for o in out] == sorted(o["ts"] for o in out)


def test_mmr_avoids_near_duplicates():
    chat = "".join(f"01/10/2026, 1{i}:00 - Amit: Razorpay payment integration completed and tested\n" for i in range(4)) + \
           "01/10/2026, 15:00 - Karan: Found a bug, cart page crashes on submit\n"
    out = summarize_messages(parse(chat), k=2)
    assert {"Amit", "Karan"} == {o["sender"] for o in out}


def test_period_filter(msgs):
    out = summarize_messages(msgs, k=10, date_from=msgs[-3].ts)
    assert all(o["ts"] >= msgs[-3].ts for o in out)


# ------------------------------------------------------------- executive summary
def test_summary_grounded(result):
    s = build_summary(result)
    assert verify_summary(s, result) == []
    kinds = [c["kind"] for c in s["claims"]]
    assert kinds[0] == "progress" and "health" in kinds and "risk" in kinds and "action" in kinds
    assert f"{result['progress']['overall']:g}%" in s["text"]
    assert "1 potential contradiction needs verification" in s["text"]


def test_summary_verification_catches_tampering(result):
    s = build_summary(result)
    s["claims"][0]["facts"]["overall"] = 99.0
    s["claims"][0]["task_ids"].append("task_does_not_exist")
    probs = verify_summary(s, result)
    assert any("progress mismatch" in p for p in probs) and any("unknown task id" in p for p in probs)


def test_summary_no_tasks():
    r = analyze(parse("01/10/2026, 10:00 - A: good morning\n01/10/2026, 10:05 - B: hi\n"))
    s = build_summary(r)
    assert "cannot be calculated" in s["text"] and verify_summary(s, r) == []


def test_rouge():
    assert rouge("project is 60 percent complete", "project is 60 percent complete")["rouge1"] == 1.0
    assert rouge("a b c", "x y z") == {"rouge1": 0.0, "rouge2": 0.0, "rougeL": 0.0}
    assert 0 < rouge("the api is done", "the api is not done")["rougeL"] < 1


# ------------------------------------------------------------ daily / weekly
def test_daily_report(result, msgs):
    d = daily_report(result, msgs, date(2026, 10, 2))
    assert d["counts"]["completed"] == 1 and d["completed"][0]["title"] == "Razorpay integration"
    assert d["progress_change"] == d["progress"] - 0.0 or d["progress_change"] is not None
    assert d["bugs_fixed"] == 1 and d["bugs_reported"] == 1
    assert d["completed"][0]["evidence_message_ids"] and d["key_messages"]
    assert daily_report(result, msgs, date(2026, 10, 4))["messages"] == 0


def test_weekly_report_monday_to_sunday(result, msgs):
    start, end = week_bounds(date(2026, 10, 2))
    assert (start.weekday(), end.weekday()) == (0, 6)
    w = weekly_report(result, msgs, date(2026, 10, 2))
    assert w["week_start"] == "2026-09-28" and len(w["days"]) == 7
    assert w["this_week"]["completed"] == 1 and w["previous_week"]["completed"] == 0
    assert w["comparison"]["completed"] == 1 and w["progress_end"] == result["progress"]["overall"]
    assert sum(d["messages"] for d in w["days"]) == w["this_week"]["messages"]


# ---------------------------------------------------------------- full report
def test_full_report_has_all_28_sections(report):
    numbered = sorted(int(k.split("_")[0].rstrip("b")) for k in report if k[0].isdigit() and not k.endswith("b") or k.startswith("23b"))
    assert set(range(1, 29)) <= set(numbered)
    assert report["meta"]["project"] == "E-Commerce Website" and report["meta"]["company"] == "Acme Tech"
    assert report["2_project_overview"]["project_manager"] == "Rahul"
    assert [t["title"] for t in report["9_completed_work"]] == ["Razorpay integration"]
    assert report["13_client_pending"] and report["14_sales_pending"][0]["requested_by"] == "Sales - Neha"
    assert report["26_bottlenecks"]["blocked_tasks"] == ["login screen"]
    assert report["27_recommendations"] and report["28_next_actions"]


def test_timeline_from_real_messages(result, msgs):
    tl = build_timeline(result, msgs)
    assert tl[0]["type"] == "START" and [e["at"] for e in tl] == sorted(e["at"] for e in tl)
    ids = {m.id for m in msgs}
    assert all(set(e["source_message_ids"]) <= ids for e in tl)
    assert any(e["type"] == "COMPLETED" and "Razorpay" in e["title"] for e in tl)
    assert any(e["type"] == "REQUIREMENT_CHANGE" for e in tl)


# -------------------------------------------------------------------- exports
def test_json_export_roundtrip(report):
    data = json.loads(export.to_json(report))
    assert data["3_progress"]["overall"] == report["3_progress"]["overall"]


def test_csv_export_has_evidence_column(report):
    csv_text = export.to_csv(report, "Tasks")
    head = csv_text.splitlines()[0]
    assert "Source messages" in head and "Razorpay integration" in csv_text
    assert "Condition" not in export.to_csv(report, "Risks")


def test_xlsx_export(report):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(export.to_xlsx(report)))
    assert {"Summary", "Tasks", "Client pending", "Risks", "Timeline", "Next actions"} <= set(wb.sheetnames)
    tasks = wb["Tasks"]
    assert tasks["A1"].value == "Task" and tasks.max_row == 6  # header + 5 tasks
    assert "Project is currently" in " ".join(str(c.value) for c in wb["Summary"]["B"] if c.value)


def test_pdf_export_is_valid(report):
    pdf = export.to_pdf(report)
    assert pdf.startswith(b"%PDF") and len(pdf) > 3000


def test_pdf_escapes_markup():
    r = analyze(parse("01/10/2026, 10:00 - Rahul: Amit please build the <b>invoice</b> & export by Friday\n"
                      "02/10/2026, 10:00 - Amit: invoice export done\n"))
    rep = full_report(r, parse("01/10/2026, 10:00 - Rahul: x\n"), {"name": "A & B <test>", "company": "C"})
    assert export.to_pdf(rep).startswith(b"%PDF")


# ------------------------------------------------------------------------- API
def test_report_and_summary_endpoints(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    import app.api.main as main
    c = TestClient(main.app)
    f = {"file": ("c.txt", io.BytesIO(CHAT.encode()))}
    s = c.post("/ai/summarize", files=f).json()
    assert s["problems"] == [] and "Project is currently" in s["summary"]["text"] and s["key_messages"]
    for fmt, magic in (("pdf", b"%PDF"), ("xlsx", b"PK"), ("json", b"{"), ("csv", b"Task")):
        r = c.post(f"/ai/report?format={fmt}&project=Shop&client=Sharma", files={"file": ("c.txt", io.BytesIO(CHAT.encode()))})
        assert r.status_code == 200 and r.content.startswith(magic), fmt
    assert c.post("/ai/report?format=exe", files={"file": ("c.txt", io.BytesIO(CHAT.encode()))}).status_code == 422
